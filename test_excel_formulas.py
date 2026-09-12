from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook

from bill_extractor.export import export_excel
from bill_extractor.models import BillRecord


def test_formula_export_preserves_inputs_and_groups_customers():
    records = [BillRecord({"customer_id": customer, "kwh_consumed": units,
                           "meter_number": "=untrusted", "total_payable": 123},
                          "UGVCL", "bill.pdf", [1])
               for customer, units in [("a", 100), ("b", 0), ("a", 200)]]
    original = deepcopy(records)
    workbook = load_workbook(BytesIO(export_excel(records)))
    sheet = workbook["Extracted Bills"]
    assert records == original
    assert sheet["B2"].data_type == "s"
    assert sheet["J2"].value == "-"
    assert sheet["J3"].value == "-"
    assert sheet["J4"].value == '=IF(AND(ISNUMBER(I4),ISNUMBER(I2),I2<>0),(I4-I2)/I2,"-")'
    assert sheet["J4"].number_format == "0.00%"
    # Sample references must follow the reordered adjustment/payable columns.
    assert sheet["BE2"].value == '=IF(COUNT(AR2,AS2,AT2,AU2,AV2,AW2,AX2,AY2,BB2,BC2,AZ2)>0,SUM(AR2,AS2,AT2,AU2,AV2,AW2,AX2,AY2,BB2,BC2,AZ2),"-")'
    assert sheet["BI2"].value == '=IF(AND(ISNUMBER(BG2)),SUM(BG2,BH2),"-")'
    assert sheet["BK2"].value == '=IF(AND(ISNUMBER(AN2),ISNUMBER(I2),I2<>0),AN2/I2,"-")'
    summary = workbook["Summary"]
    assert summary["I2"].value == '=IF(COUNT(\'Extracted Bills\'!I2,\'Extracted Bills\'!I4)>0,AVERAGE(\'Extracted Bills\'!I2,\'Extracted Bills\'!I4),"-")'
    assert "SUM(" in summary["I3"].value
    assert "I3" in summary["I4"].value
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad
    assert sheet.auto_filter.ref == "A1:BM4"


def test_empty_export_keeps_headers():
    workbook = load_workbook(BytesIO(export_excel([])))
    assert workbook["Extracted Bills"].max_row == 1
    assert "Summary" not in workbook.sheetnames
