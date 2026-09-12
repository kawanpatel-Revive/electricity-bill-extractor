"""Native Excel and JSON exports."""

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bill_extractor.models import BillRecord
from bill_extractor.formula_cache import save_with_formula_results
from bill_extractor.schema import FIELDS
from bill_extractor.excel_formulas import COLUMNS, SUMMARY_AVERAGES, SUMMARY_SUMS, TOTAL_SUMS, row_formulas


def export_excel(records: list[BillRecord]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extracted Bills"
    sheet.append([field.label for field in FIELDS])

    customer_rows: dict[str, list[int]] = {}
    for row, record in enumerate(records, 2):
        sheet.append([record.values.get(field.key) if record.values.get(field.key) is not None else "-" for field in FIELDS])
        for cell in sheet[row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        customer = str(record.values.get("customer_id") or record.filename)
        prior_rows = customer_rows.setdefault(customer, [])
        for key, formula in row_formulas(row, prior_rows[-1] if prior_rows else None).items():
            sheet[f"{COLUMNS[key]}{row}"] = formula
        prior_rows.append(row)

    for column, field in enumerate(FIELDS, 1):
        if field.kind != "percent" and field.key != "kwh_increase_percent":
            continue
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            if isinstance(cell.value, (int, float)) or cell.data_type == "f":
                cell.number_format = "0.00%"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, field in enumerate(FIELDS, 1):
        longest = max([len(field.label)] + [len(str(record.values.get(field.key) or "-")) for record in records])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 34)

    if records:
        summary = workbook.create_sheet("Summary")
        summary.append([field.label for field in FIELDS])
        for customer, rows in customer_rows.items():
            for label, averages, sums in (("Summary (SUM / AVERAGE)", SUMMARY_AVERAGES, SUMMARY_SUMS), ("Totals", frozenset(), TOTAL_SUMS)):
                row = summary.max_row + 1
                summary.cell(row, 1, label)
                summary.cell(row, 3, customer).data_type = "s"
                for index, field in enumerate(FIELDS, 1):
                    if field.key not in averages | sums:
                        continue
                    references = ",".join(f"'Extracted Bills'!{COLUMNS[field.key]}{source_row}" for source_row in rows)
                    operation = "AVERAGE" if field.key in averages else "SUM"
                    cell = summary.cell(row, index, f'=IF(COUNT({references})>0,{operation}({references}),"-")')
                    if field.kind == "percent":
                        cell.number_format = "0.00%"
        for cell in summary[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in sheet.column_dimensions:
            summary.column_dimensions[column].width = sheet.column_dimensions[column].width
        summary.freeze_panes = "D2"

    details = workbook.create_sheet("Extraction Details")
    details.append(["Provider", "Filename", "Pages", "Warnings"])
    for record in records:
        details.append([record.provider, record.filename, ", ".join(map(str, record.pages)), " | ".join(record.warnings) or "-"])
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    details.freeze_panes = "A2"
    details.column_dimensions["A"].width = 20
    details.column_dimensions["B"].width = 45
    details.column_dimensions["C"].width = 15
    details.column_dimensions["D"].width = 80

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return save_with_formula_results(workbook)


def export_json(records: list[BillRecord]) -> str:
    return json.dumps([record.as_dict() for record in records], indent=2, ensure_ascii=False)
