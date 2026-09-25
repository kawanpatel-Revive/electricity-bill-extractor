from pathlib import Path
from copy import deepcopy

from openpyxl import load_workbook
import pytest

from bill_extractor.export import export_excel
from bill_extractor.document import extract_pages
from bill_extractor.providers.ugvcl import UGVCLParser
from bill_extractor.presentation import display_dataframe
from bill_extractor.schema import FIELDS
from bill_extractor.service import InputFile, extract_files


DATA = Path("data")


def extract(filename: str):
    path = DATA / filename
    return extract_files([InputFile(path.name, path.read_bytes())], use_ocr=False)


def test_schema_matches_reference_workbook():
    workbook = load_workbook(DATA / "Angiplast_Bill_Analysis.xlsx", read_only=True)
    expected = [cell.value for cell in workbook["Angiplast"][1]][:58]

    fuel_percent_index = expected.index("Fuel Surcharge %") + 1
    expected[fuel_percent_index:fuel_percent_index] = ["Base FPPAS", "FPPAS Charges", "FPPAS %"]
    previous_dues_index = expected.index("Previous Dues") + 1
    expected[previous_dues_index:previous_dues_index] = [
        "Electricity Duty Credits",
        "TOU Charge Credits",
        "TDS Credits",
        "Other Debits",
    ]
    expected.remove("Security Deposit Interest")
    expected.remove("Other Credits")
    advance_index = expected.index("Adv Payment/ Adjustment")
    expected[advance_index:advance_index] = ["Security Deposit Interest", "Other Credits"]
    unit_rate_index = expected.index("Unit Rate (Total Consumption Charge)")
    demand_unit_rate_index = expected.index("Total Consumption-Demand Charge Unit Rate\n\n")
    expected[unit_rate_index], expected[demand_unit_rate_index] = (
        expected[demand_unit_rate_index],
        expected[unit_rate_index],
    )

    assert len(FIELDS) == 65
    assert expected == [field.label for field in FIELDS]


def test_ugvcl_merged_bill_matches_reference_readings():
    records = extract("Angiplast_UGVCL-INVOICES 2025-26 Merge.pdf")
    by_month = {record.values["billing_month"]: record for record in records}
    workbook = load_workbook(DATA / "Angiplast_Bill_Analysis.xlsx", data_only=True, read_only=True)
    sheet = workbook["Angiplast"]

    assert len(records) == 12
    for row in range(2, 14):
        month = sheet.cell(row, 1).value
        assert by_month[month].values["kwh_consumed"] == sheet.cell(row, 9).value
        assert by_month[month].values["total_consumption_charges"] == pytest.approx(sheet.cell(row, 38).value)

    april = by_month["APR-2025"].values
    assert april["customer_id"] == "65500"
    assert april["meter_number"] == "GHBD1806"
    assert april["solar_generation_units"] == 37920
    assert april["solar_setoff_units"] == 423
    assert april["solar_export_units"] == 5003
    assert april["solar_banking_units"] == 32917
    assert april["calculated_adjustment"] == 22648.10
    assert april["total_payable"] == 424939.03
    assert april["net_less_demand_unit_rate"] == pytest.approx(
        (april["total_payable"] - april["demand_charges"]) / april["kwh_consumed"]
    )

    september = by_month["SEP-2025"]
    assert september.values["other_debits"] == pytest.approx(13.35)
    assert september.values["calculated_adjustment"] == pytest.approx(
        september.values["advance_adjustment"]
    )
    assert not september.warnings

    february = by_month["FEB-2026"].values
    assert february["solar_setoff_units"] == 152
    assert february["solar_net_billed_units"] == 64864
    assert february["solar_setoff_credit"] == pytest.approx(-957.00)
    assert february["other_credits"] == 0

    march = by_month["MAR-2026"].values
    assert march["solar_setoff_units"] == 83
    assert march["solar_net_billed_units"] == 54889
    assert march["solar_setoff_credit"] == pytest.approx(-522.00)
    assert march["other_credits"] == 0


def test_ugvcl_combines_multirow_adjustments():
    may = {record.values["billing_month"]: record for record in extract("Angiplast_UGVCL-INVOICES 2025-26 Merge.pdf")}["MAY-2025"].values

    assert may["solar_setoff_units"] == 2300
    assert may["solar_export_units"] == 18318
    assert may["solar_banking_units"] == 71442
    assert may["solar_generation_units"] == 89760
    assert may["security_deposit_interest"] == -115233.75
    assert may["other_credits"] == -975.24


def test_torrent_power_parser():
    record = extract("Torrent Bill Marking_260810_174038.pdf")[0]
    values = record.values

    assert record.provider == "Torrent Power"
    assert values["billing_month"] == "JAN-2026"
    assert values["customer_id"] == "100358213"
    assert values["contract_demand"] == 900
    assert values["actual_max_demand"] == 612
    assert values["billing_demand"] == 765
    assert values["kwh_consumed"] == 249690
    assert values["fuel_surcharge"] is None
    assert values["base_fppas"] == 926797.08
    assert values["fppas_charges"] == 76815.50
    assert values["fppas_percent"] == pytest.approx(0.034)
    assert values["electricity_duty"] == 359765.81
    assert values["solar_generation_units"] == 37805
    assert values["solar_net_billed_units"] == 249139
    assert values["total_payable"] == 2779320.78
    assert not record.warnings


def test_torrent_old_fpppa_layout_populates_charge_components():
    values = extract("MARCH 25.pdf")[0].values

    assert values["demand_charges"] == 198900.00
    assert values["energy_charges"] == 992204.85
    assert values["base_fppas"] == 830835.27
    assert values["fppas_charges"] is None
    assert values["fuel_surcharge"] is None
    assert values["total_energy_charges"] == pytest.approx(2079092.99)
    assert values["total_consumption_charges"] == pytest.approx(2079092.99)


def test_torrent_combined_credit_is_split_by_solar_note():
    values = extract("DEC 25.pdf")[0].values

    assert values["solar_credit"] == pytest.approx(-9911.25)
    assert values["other_credits"] == pytest.approx(-2906.11)
    assert values["calculated_adjustment"] == pytest.approx(25428.54)


def test_torrent_security_interest_and_tds_are_separated():
    values = extract("May'25.pdf")[0].values

    assert values["security_deposit_interest"] == pytest.approx(-137642.23)
    assert values["other_debits"] == pytest.approx(13765.00)


def test_torrent_ocr_layout_recovers_header_setoff_and_banking_units():
    text = """
Torrent Power
ACCUMAX LAB DEVICES PVT LTD CONTRACT DEMAND BILLING MONTH HTMD1
PLOT NO.14,15,16 and 32 900 KW July 2025
GIDC BILLING DEMAND READING DATE CUSTOMERID
765.0 KW 31/07/25 100358213
Registered Mobile: *5552 100 01/08/25
Meter No.:29800060
Units 763.000 334410.000 122480.000 335970.000 94160.000
Energy charges (A) 1517692.10
Fixed demand charges (B) 198900.00
Excess demand charges 0.00
Base FPPAS @Rs. 3.72/unit (C) 1243142.16
FPPAS charges @3.40% of (A+B+C) 100630.96
TOU charges 122480.00
Power Factor adjustment charges 7022.61-
NTC rebate 28248.00-
Total energy charges 3147574.61
Total government duty @15.00% 472136.19
Banking charges (Solar generation unit-Excess solar unit)@Rs.1.10 43331.20
Other debit 0.00
Credit 3748.5-
Previous dues 3459.39-
Amount due 3655834.11
Solar generation units are: 41058, Net billed units- 334178
Solar Sstoff Units 232.00
Credit of Rs. 3748.50 for 1666 excess Solar
"""

    from bill_extractor.providers.torrent import TorrentParser

    values = TorrentParser().parse(text)

    assert values["billing_month"] == "JUL-2025"
    assert values["tariff_category"] == "HTMD1"
    assert values["solar_setoff_units"] == 232
    assert values["solar_banking_units"] == 39392
    assert values["solar_export_units"] == 1666
    assert values["electricity_duty"] == pytest.approx(472136.19)
    assert values["fppas_charges"] == pytest.approx(100630.96)


@pytest.mark.parametrize(
    ("filename", "amount_due"),
    [
        ("JAN 26.pdf", 2779320.78),
        ("FEB 26.pdf", 2858661.96),
    ],
)
def test_torrent_uses_precise_amount_due_and_excludes_demand_from_bd(filename, amount_due):
    record = extract(filename)[0]
    values = record.values

    assert values["total_payable"] == pytest.approx(amount_due)
    assert values["consumption_demand_unit_rate"] == pytest.approx(
        (values["total_consumption_charges"] - values["demand_charges"])
        / values["kwh_consumed"]
    )


def test_excel_export_uses_reference_headers_and_missing_marker():
    records = extract_files(
        [InputFile("S P METAL PGVCL.pdf", Path("S P METAL PGVCL.pdf").read_bytes())],
        use_ocr=False,
    )
    workbook = load_workbook(filename=__import__("io").BytesIO(export_excel(records)), data_only=False)
    sheet = workbook["Extracted Bills"]

    assert [cell.value for cell in sheet[1]] == [field.label for field in FIELDS]
    tariff_column = next(index for index, field in enumerate(FIELDS, 1) if field.key == "tariff_category")
    solar_column = next(index for index, field in enumerate(FIELDS, 1) if field.key == "solar_generation_units")
    assert sheet.cell(2, tariff_column).value == "LTMD"
    assert sheet.cell(2, solar_column).value == '=IF(COUNT(U2,W2)>0,SUM(U2,W2),"-")'


def test_percentage_columns_use_excel_percentage_format():
    records = extract("Angiplast_UGVCL-INVOICES 2025-26 Merge.pdf")
    workbook = load_workbook(filename=__import__("io").BytesIO(export_excel(records)), data_only=False)
    sheet = workbook["Extracted Bills"]
    solar_export_column = next(
        index for index, field in enumerate(FIELDS, 1) if field.key == "solar_export_percent"
    )

    assert sheet.cell(2, solar_export_column).value == '=IF(AND(ISNUMBER(U2),ISNUMBER(Q2),Q2<>0),U2/Q2,"-")'
    assert sheet.cell(2, solar_export_column).number_format == "0.00%"


def test_streamlit_preview_is_arrow_safe_with_mixed_missing_and_numeric_values():
    records = extract_files(
        [InputFile("S P METAL PGVCL.pdf", Path("S P METAL PGVCL.pdf").read_bytes())],
        use_ocr=False,
    )
    records.append(deepcopy(records[0]))
    records[1].values["kwh_increase_percent"] = 0.125

    frame = display_dataframe(records)

    assert all(str(dtype) == "string" for dtype in frame.dtypes)
    assert frame["% Increase in kWh"].tolist() == ["-", "0.125"]
    assert "% of Total units (Night)" in frame.columns
    assert "% of Total units (TOU)" in frame.columns
    assert frame.columns.is_unique

    pyarrow = pytest.importorskip("pyarrow")
    pyarrow.Table.from_pandas(frame)


def test_sparse_selectable_text_page_does_not_trigger_ocr():
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "This is a system generated bill. Hence no signature required.")
    data = document.tobytes()
    document.close()

    pages, warnings = extract_pages(data, "footer.pdf", use_ocr=True)

    assert pages[0].text.strip() == "This is a system generated bill. Hence no signature required."
    assert not pages[0].used_ocr
    assert not warnings


def test_dgvcl_parser_extracts_ht_bill_tables():
    text = """
Dakshin Gujarat Vij Company Ltd. DGVCL
HT BILL FOR THE MONTH OF :AUG-2025
Consumer No: Tarrif Demand Demand Demand Demand DMD
12459 HTP-1 200 170 149 170 1154037 0.00
Meter No: Make CTPT Make CTPT Srno CT Ratio PT Ratio Normal
DGHTS215 SECURE KVAH KVARH 3
A.Total Units B.Night Units C.TOU D.1/3 Of Units in A E.Night Concession Units
56363 15054 21836 18788 15054
Tot Demand 170 25500
Energy Charges 56363 4 225452.00
Fuel charge 56363 2.45 138089.35
PF Rebate 225452 -2.45% -5523.57
EHV Rebate 225452.00 1.00 -2254.52
TOU 21836 0.45 9826.20
SUMMARY OF CHARGES
Demand Charge Energy Charge Fuel PF Adj/Rebate Night Rebate EHV Time Of Use GT Charges Tot Consumption Charge
25500.00 225452.00 138089.35 -5523.57 0.00 -2254.52 9826.20 0.00 391089.46
Electricity Duty Meter Charges Cross Subsidy Wheeling Charges Parallel Operation Charges Current Month's Bill Outstanding Arrears
58663.42 0.00 449752.88 0.10
Charges Delayed Payment Adjust. Adv.Payment / Net Payable TCS Total Payable PREV.BILL TCS Cr Reading Date
0.00 -19305.86 430447.12 0.00 430447.12 0.00 16-08-2025
"""

    from bill_extractor.providers.dgvcl import DGVCLParser

    parser = DGVCLParser()
    values = parser.parse(text)

    assert parser.matches(text)
    assert values["billing_month"] == "AUG-2025"
    assert values["customer_id"] == "12459"
    assert values["meter_number"] == "DGHTS215"
    assert values["contract_demand"] == 200
    assert values["actual_max_demand"] == 149
    assert values["billing_demand"] == 170
    assert values["kwh_consumed"] == 56363
    assert values["demand_charges"] == 25500
    assert values["total_consumption_charges"] == pytest.approx(391089.46)
    assert values["electricity_duty"] == pytest.approx(58663.42)
    assert values["current_month_bill"] == pytest.approx(449752.88)
    assert values["outstanding_arrears"] == pytest.approx(0.10)
    assert values["advance_adjustment"] == pytest.approx(-19305.86)
    assert values["total_payable"] == pytest.approx(430447.12)


def test_dgvcl_parser_handles_split_summary_rows_and_ocr_variants():
    text = """
Dakshin Gujarat Vij Company Ltd. DGVCL
HT BILL, FOR THE MONTH OF :JUI-2025
Consumer No: Tarril Contract 85% Contract Actunl Max. Billing Excess Cont.
Demand Demand Vemand Demand DMD
12459 HTPI 200 170 134 170 1154037 0.00
Supp Voltage KWH KVAH KVARH Avg PF MF Actual Max DMD during day
11 52788 52793 198 .999 3
A.Total Units B.Night Units C.TOU D.1/3 Of Units in A E.Night Concession Units F.Connection G.Consumer Type
Date
52788 14034 20736 17596 14034 20-07-2022
SUMMARY OF CHARGES
Demand Charge Energy Charge Fuel Surcharge PF Adj/Rebate Night Rebate Rebate EHV Time Of Use GT Charges Tot Consumption Charge
25500.00 211152.00 129330.60 -5173.22 0.00 -2111.52 9331.20 0.00 368029.06
Electricity Duty Meter Charges Cross Subsidy Wheeling Charges Parallel Operation Charges Current Outstanding Arrears
55204.36 0.00 MOnth's Bill
Delayed Payment Adv.Payment / 423233.42 0.68
Charges Adjust. Net Payable TCS Total Payable PREV.BULL TCS Cr Reading Date Bill Date
0.00 0.00 423234.10 0.00 423234.10 0.00 16-07-2025
"""

    from bill_extractor.providers.dgvcl import DGVCLParser

    values = DGVCLParser().parse(text)

    assert values["billing_month"] == "JUL-2025"
    assert values["tariff_category"] == "HTP-1"
    assert values["average_power_factor"] == pytest.approx(0.999)
    assert values["night_units"] == 14034
    assert values["tou_kwh"] == 20736
    assert values["total_consumption_charges"] == pytest.approx(368029.06)
    assert values["electricity_duty"] == pytest.approx(55204.36)
    assert values["current_month_bill"] == pytest.approx(423233.42)
    assert values["outstanding_arrears"] == pytest.approx(0.68)
    assert values["advance_adjustment"] == 0
    assert values["total_payable"] == pytest.approx(423234.10)


def test_dgvcl_parser_recovers_tou_charge_from_summary():
    text = """
Dakshin Gujarat Vij Company Ltd. DGVCL
HT BILL FOR THE MONTH OF :JUN-2025
SUMMARY OF CHARGES
Demand Charge Energy Charge Fuel Surcharge PF Adj/Rebate Night Rebate EHV Rebate Time Of Use Charges GT Charges Tot Consumption Charge
25500.00 245616.00 150439.80 -6017.59 0.00 -2456.16 10464.75 0.00 423546.80
Electricity Duty Meter Charges Current Month's Bill Outstanding Arrears
63532.02 0.00 487078.82 0.86
Delayed Payment Adv.Payment / Net Payable TCS Total Payable Reading Date
0.00 0.00 487079.68 0.00 487079.68 0.00 16-06-2025
"""

    from bill_extractor.providers.dgvcl import DGVCLParser

    values = DGVCLParser().parse(text)

    assert values["tou_charges"] == pytest.approx(10464.75)
    assert values["total_consumption_charges"] == pytest.approx(423546.80)


def test_ugvcl_wrapped_meter_and_adjustments_are_parsed():
    text = """
HT BILL FOR THE MONTH OF : APR-2026
Meter No: Make CTPT Make CTPT Srno CT Ratio PT Ratio Meter Status
Meter
Constant
GHBD0736 L&T 15 Normal
Adjustment Details Report for APR-2026/td>
Credit Board
39577.20 0.00 CREDIT BC FOR CONSUMPTION BETWEEN 11 AM TO 3 PM, MONTH-MAR-26, UNIT-65962
Charges
Credit Board
234.38 35.00 SOLAR ADJ FOR THE MONTH OF -MAR-26
Charges
Credit ED
4382.97 0.00 CREDIT ED FOR CONSUMPTION BETWEEN 11 AM TO 3 PM, MONTH-MAR-26, UNIT-65962
Charges
Credit ED
25.96 0.00 SOLAR ADJ FOR THE MONTH OF -MAR-26
Charges
Credit TDS 3202.00 0.00 TDS ADJUSTMENT 3-2026
Debit Banking
14016.00 0.00 SOLAR BANKING CHARGES FOR THE MONTH - MAR-26, UNIT-9344
Charge
"""

    values = UGVCLParser().parse(text)

    assert values["meter_number"] == "GHBD0736"
    assert values["solar_setoff_units"] == 35
    assert values["solar_setoff_credit"] == pytest.approx(-234.38)
    assert values["electricity_duty_credits"] == pytest.approx(-4408.93)
    assert values["tou_charge_credits"] == pytest.approx(-39577.20)
    assert values["tds_credits"] == pytest.approx(-3202.00)
    assert values["solar_banking_units"] == 9344
    assert values["solar_banking_charges"] == 14016
    assert values["other_credits"] == 0
    assert values["security_deposit_interest"] == 0


def test_ugvcl_other_debits_reconcile_adjustment():
    text = """
Adjustment Details Report for OCT-2025/td>
Credit Board Charges 848.74 127.00 SOLAR ADJ SEP-25
Credit ED Charges 93.77 0.00 SOLAR ADJ SEP-25
Credit TDS 3625.00 0.00 TDS ADJUSTMENT 9-2025
Debit Banking Charge 12012.00 0.00 SOLAR BANKING CHARGE SEP-25,U-8008
Debit Electricity Duty 0.05 0.00 ED RECOVERY AGAINST FC
Debit Fuel Surcharge 0.45 0.00 FC RECOVERY IN SOLAR SET OFF FOR THE MONTH OF JULY-25
"""

    values = UGVCLParser().parse(text)

    assert values["solar_setoff_units"] == 127
    assert values["solar_setoff_credit"] == pytest.approx(-848.74)
    assert values["electricity_duty_credits"] == pytest.approx(-93.77)
    assert values["tds_credits"] == pytest.approx(-3625.00)
    assert values["solar_banking_units"] == 8008
    assert values["other_debits"] == pytest.approx(0.50)
    assert values["other_credits"] == 0


@pytest.mark.parametrize("description", ["Cedit", "Credit", "Debit"])
def test_ugvcl_wrapped_fuel_surcharge_adjustments(description):
    text = f"""
{description} Fuel
40379.58 0.00 Fuel Surcharge for Jul-2025
Surcharge
{description} Fuel
97749.00 0.00 Fuel Surcharge for Aug-2025
Surcharge
"""
    values = UGVCLParser().parse(text)

    assert values["other_credits"] == pytest.approx(
        0.0 if description == "Debit" else -138128.58
    )
    assert values["other_debits"] == pytest.approx(
        138128.58 if description == "Debit" else 0.0
    )
    assert values["fuel_surcharge"] is None


def test_medha_ugvcl_adjustment_comment_fixes():
    data_dir = DATA / "MEDHA UGVCL"
    filenames = [
        "MEDHA UGVCL-APRIL 25.pdf",
        "MEDHA UGVCL-MAY 25.pdf",
        "MEDHA UGVCL-JUNE 25.pdf",
        "MEDHA UGVCL-JULY 25.pdf",
        "MEDHA UGVCL-AUG 25.pdf",
        "MEDHA UGVCL-SEP 25.pdf",
        "MEDHA UGVCL-OCT 25.pdf",
        "MEDHA UGVCL NOV 25.pdf",
        "MEDHA UGVCL-DEC 25.pdf",
        "MEDHA UGVCL-JAN 26.pdf",
        "MEDHA UGVCL-FEB 26.pdf",
    ]
    records = extract_files(
        [InputFile(name, (data_dir / name).read_bytes()) for name in filenames],
        use_ocr=False,
    )
    by_month = {record.values["billing_month"]: record for record in records}

    assert len(records) == 11
    expected_solar = {
        "APR-2025": (1855, 5474, 64533, -12140.98),
        "MAY-2025": (1355, 5178, 75571, -8811.57),
        "JUN-2025": (480, 2058, 68035, -3121.44),
        "JUL-2025": (59, 0, 55584, -383.68),
        "AUG-2025": (543, 736, 45453, -3531.13),
        "SEP-2025": (1425, 690, 42557, -9053.03),
        "OCT-2025": (1351, 1723, 45774, -8582.90),
        "NOV-2025": (1435, 515, 39513, -9116.56),
        "DEC-2025": (708, 4532, 31420, -4497.92),
        "JAN-2026": (1655, 1317, 29383, -10514.22),
        "FEB-2026": (636, 1335, 36993, -4040.51),
    }
    for month, (setoff, export, banking, setoff_credit) in expected_solar.items():
        values = by_month[month].values
        assert values["solar_setoff_units"] == setoff
        assert values["solar_net_billed_units"] == values["kwh_consumed"] - setoff
        assert values["solar_export_units"] == export
        assert values["solar_banking_units"] == banking
        assert values["solar_generation_units"] == export + banking
        assert values["solar_setoff_credit"] == pytest.approx(setoff_credit)
        assert values["previous_dues"] == 0.0

    july = by_month["JUL-2025"].values
    assert july["tds_credits"] == pytest.approx(-4855.00)
    assert july["other_credits"] == pytest.approx(-4801.00)

    september = by_month["SEP-2025"].values
    assert september["other_debits"] == pytest.approx(38.10)
    assert september["other_credits"] == pytest.approx(-138128.58)
    assert september["calculated_adjustment"] == pytest.approx(-89567.04)
    assert september["calculated_adjustment"] == pytest.approx(september["advance_adjustment"])

    february = by_month["FEB-2026"].values
    assert february["demand_charges"] == pytest.approx(549207.14285714)
    assert february["total_energy_charges"] == pytest.approx(3777701.31)
    assert february["total_consumption_charges"] == pytest.approx(3839415.60)


def test_photographed_bill_and_adjustment_are_merged():
    filenames = ["1000371676.jpg", "1000371677.jpg"]
    records = extract_files(
        [InputFile(name, (DATA / name).read_bytes()) for name in filenames],
        use_ocr=True,
    )

    assert len(records) == 1
    record = records[0]
    values = record.values
    assert record.provider == "UGVCL"
    assert values["billing_month"] == "MAY-2025"
    assert values["customer_id"] == "65500"
    assert values["tariff_category"] == "HTP-I"
    assert values["kwh_consumed"] == 52176
    assert values["solar_generation_units"] == 89760
    assert values["solar_setoff_units"] == 2300
    assert values["solar_export_units"] == 18318
    assert values["solar_banking_units"] == 71442
    assert values["net_payable"] == 295385.72
    assert values["total_payable"] == 295385.72
    assert not record.warnings


@pytest.mark.parametrize(
    (
        "filename",
        "setoff",
        "export",
        "banking",
        "banking_charge",
        "solar_credit",
        "setoff_credit",
        "electricity_duty_credit",
        "tds_credit",
        "other_credit",
    ),
    [
        ("1_APRIL-2025.pdf", 528, 202, 35318, 52977.00, -333.30, -3405.60, -510.84, -1108.00, 0.00),
        ("2_MAY-2025.pdf", 820, 1425, 53095, 79642.50, -2351.25, -5315.24, -1159.17, 0.00, -2412.54),
        ("3_JUNE-2025.pdf", 869, 1613, 51788, 77682.00, -2661.45, -5570.29, -835.54, 0.00, 0.00),
        ("4_JULY-2025.pdf", 1137, 719, 55333, 82999.50, -1186.35, -7219.95, -1082.99, -4177.00, 0.00),
        ("5_AUG-2025.pdf", 1731, 607, 27986, 41979.00, -1001.55, -10960.69, -1644.10, 0.00, 0.00),
    ],
)
def test_s21_ugvcl_adjustments(
    filename,
    setoff,
    export,
    banking,
    banking_charge,
    solar_credit,
    setoff_credit,
    electricity_duty_credit,
    tds_credit,
    other_credit,
):
    record = extract(filename)[0]
    values = record.values

    assert values["solar_setoff_units"] == setoff
    assert values["solar_export_units"] == export
    assert values["solar_banking_units"] == banking
    assert values["solar_generation_units"] == export + banking
    assert values["solar_net_billed_units"] == values["kwh_consumed"] - setoff
    assert values["solar_banking_charges"] == pytest.approx(banking_charge)
    assert values["solar_credit"] == pytest.approx(solar_credit)
    assert values["solar_setoff_credit"] == pytest.approx(setoff_credit)
    assert values["electricity_duty_credits"] == pytest.approx(electricity_duty_credit)
    assert values["tds_credits"] == pytest.approx(tds_credit)
    assert values["other_credits"] == pytest.approx(other_credit)
    assert values["calculated_adjustment"] == pytest.approx(values["advance_adjustment"])
    assert values["net_less_demand_unit_rate"] == pytest.approx(
        (values["total_payable"] - values["demand_charges"])
        / values["kwh_consumed"]
    )
    assert not record.warnings
